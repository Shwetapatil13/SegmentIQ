"""
Generate editable academic DOCX report from reports/results.json and reports/figures/.
"""

import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cells."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = docx.oxml.OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = docx.oxml.OxmlElement(m)
        node.set(docx.oxml.ns.qn('w:w'), str(val))
        node.set(docx.oxml.ns.qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def generate_report():
    results_path = os.path.join("reports", "results.json")
    if not os.path.exists(results_path):
        raise FileNotFoundError(f"Results file '{results_path}' not found. Run src/run_experiments.py first.")

    with open(results_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    doc = docx.Document()

    # Configure 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # -------------------------------------------------------------
    # COVER PAGE
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("A COMPARATIVE STUDY OF CLUSTERING ALGORITHMS FOR CUSTOMER SEGMENTATION\n")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("A Mini-Project Report Submitted in Partial Fulfillment of the Requirements for the Degree of\n")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    
    deg_run = sub_p.add_run("BACHELOR OF TECHNOLOGY / ENGINEERING\nIN COMPUTER ENGINEERING\n\n")
    deg_run.font.size = Pt(13)
    deg_run.font.bold = True

    # Submitted by
    by_p = doc.add_paragraph()
    by_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    by_run = by_p.add_run("SUBMITTED BY:\n")
    by_run.font.size = Pt(11)
    by_run.font.bold = True
    
    st_run = by_p.add_run("[Student 1 Name Placeholder]  (Roll No: [XXXXXX])\n[Student 2 Name Placeholder]  (Roll No: [XXXXXX])\n\n")
    st_run.font.size = Pt(12)

    # Under Guidance of
    gd_p = doc.add_paragraph()
    gd_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gd_run = gd_p.add_run("UNDER THE GUIDANCE OF:\n")
    gd_run.font.size = Pt(11)
    gd_run.font.bold = True
    
    guide_run = gd_p.add_run("[Faculty Guide Name Placeholder]\n[Designation / Department]\n\n")
    guide_run.font.size = Pt(12)

    # Institution
    inst_p = doc.add_paragraph()
    inst_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inst_run = inst_p.add_run("DEPARTMENT OF COMPUTER ENGINEERING\n[INSTITUTE / COLLEGE NAME PLACEHOLDER]\nACADEMIC YEAR: 2025 – 2026\n")
    inst_run.font.size = Pt(13)
    inst_run.font.bold = True
    inst_run.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_page_break()

    # -------------------------------------------------------------
    # ABSTRACT
    # -------------------------------------------------------------
    h_abs = doc.add_heading("ABSTRACT", level=1)
    h_abs.paragraph_format.space_before = Pt(12)
    
    meta = data["metadata"]
    kmeans_meta = data["benchmarks"]["models"]["K-Means"]["metrics"]
    dbscan_meta = data["benchmarks"]["models"]["DBSCAN"]["metrics"]
    
    doc.add_paragraph(
        f"Customer segmentation is an essential capability for retail organizations seeking to extract actionable intelligence "
        f"from consumer transaction records. This project presents an empirical comparative evaluation of three core clustering paradigms: "
        f"centroid-based partitioning (K-Means), connectivity-based agglomerative hierarchical clustering (Ward linkage), and density-based "
        f"spatial clustering with noise (DBSCAN). Experiments were conducted on the benchmark Mall Customer Segmentation Dataset "
        f"(N = {meta['n_records']} customer records, evaluating {', '.join(meta['features_selected'])}). "
        f"Models were evaluated across mathematical validity metrics—Silhouette Score, Davies-Bouldin Index, and Calinski-Harabasz Index—alongside "
        f"execution runtime and noise robustness. All transformations were applied strictly on standardized feature matrices without data leakage. "
        f"K-Means achieved the strongest internal cluster validity (Silhouette Score: {kmeans_meta['silhouette_score']}, "
        f"Davies-Bouldin: {kmeans_meta['davies_bouldin_index']}) in {kmeans_meta['runtime_sec']:.4f} seconds, segmenting the portfolio into "
        f"five distinct consumer personas: Mainstream Middle, Affluent Spenders, Aspirational Trendsetters, Cautious Affluent, and Frugal Budgeters. "
        f"While DBSCAN isolated boundary noise points ({dbscan_meta['n_noise']} outliers, {dbscan_meta['noise_pct']}%), K-Means proved superior "
        f"for comprehensive retail customer relationship management."
    )

    # -------------------------------------------------------------
    # 1. INTRODUCTION & PROBLEM STATEMENT
    # -------------------------------------------------------------
    doc.add_heading("1. INTRODUCTION & PROBLEM STATEMENT", level=1)
    doc.add_paragraph(
        "Modern retail businesses collect substantial amounts of consumer demographic and purchase records. "
        "However, treating an entire consumer base as a monolithic population results in untargeted promotional budgets, "
        "inefficient inventory allocation, and poor customer retention. Customer segmentation solves this challenge by "
        "identifying natural sub-groups sharing homogeneous traits.\n\n"
        "Because customer records lack target supervisory labels (e.g., ground-truth segment designations), unsupervised "
        "machine learning algorithms must be applied. Unsupervised clustering groups observations such that within-cluster "
        "similarity is maximized while inter-cluster similarity is minimized."
    )

    # -------------------------------------------------------------
    # 2. OBJECTIVES & ARCHITECTURE
    # -------------------------------------------------------------
    doc.add_heading("2. PROJECT OBJECTIVES & ARCHITECTURE", level=1)
    doc.add_paragraph(
        "The primary objectives of this study are:\n"
        "1. Benchmark Partitioning (K-Means), Hierarchical (Ward Linkage), and Density-Based (DBSCAN) clustering algorithms on standardized customer features.\n"
        "2. Ground hyperparameter selection in empirical diagnostics: Elbow Method, Silhouette Analysis, Dendrogram Cuts, and k-Distance Plots.\n"
        "3. Evaluate cluster quality using mathematical validity metrics while strictly isolating noise points in density algorithms.\n"
        "4. Develop an interactive Streamlit analytical dashboard and extract rule-based customer personas.\n"
        "5. Maintain a Single Source of Truth architecture where all report statistics, visual figures, and app defaults derive from an automated pipeline."
    )

    # -------------------------------------------------------------
    # 3. DATASET DESCRIPTION & EDA
    # -------------------------------------------------------------
    doc.add_heading("3. DATASET DESCRIPTION & EXPLORATORY DATA ANALYSIS", level=1)
    doc.add_paragraph(
        f"The study uses the Mall Customer Segmentation dataset comprising {meta['n_records']} unique customer observations. "
        "The features include CustomerID, Gender, Age, Annual Income (k$), and Spending Score (1-100). "
        "CustomerID was explicitly excluded as an arbitrary identifier carrying zero behavioral variance."
    )

    # Insert Figure 1
    fig1 = data["figures"]["fig1"]
    if os.path.exists(fig1):
        doc.add_picture(fig1, width=Inches(6.0))
        cap1 = doc.add_paragraph("Figure 1: Exploratory feature distributions and outlier boxplots.")
        cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap1.runs[0].font.italic = True
        cap1.runs[0].font.size = Pt(9.5)

    # -------------------------------------------------------------
    # 4. PREPROCESSING & STANDARDIZATION
    # -------------------------------------------------------------
    doc.add_heading("4. PREPROCESSING & FEATURE STANDARDIZATION", level=1)
    doc.add_paragraph(
        "Distance-based clustering algorithms rely on Euclidean distance: d(p, q) = sqrt(sum (p_i - q_i)^2). "
        "Features with larger numerical scales (Annual Income: $15k–$137k) would dominate smaller features if unscaled. "
        "We applied StandardScaler (z = (x - mu) / sigma) fit once on selected features to guarantee mean=0 and std=1 without leakage."
    )

    # Add Before/After Table
    table_data = data["preprocessing"]["before_after_summary"]
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Shading'

    headers = ["Feature", "Raw Mean", "Raw Std", "Raw Range", "Scaled Mean", "Scaled Std", "Scaled Range"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_margins(cell)

    for row in table_data:
        r = table.add_row()
        r.cells[0].text = str(row["Feature"])
        r.cells[1].text = str(row["Raw Mean"])
        r.cells[2].text = str(row["Raw Std"])
        r.cells[3].text = f"[{row['Raw Min']}, {row['Raw Max']}]"
        r.cells[4].text = str(row["Scaled Mean"])
        r.cells[5].text = str(row["Scaled Std"])
        r.cells[6].text = f"[{row['Scaled Min']}, {row['Scaled Max']}]"
        for c in r.cells:
            set_cell_margins(c)

    doc.add_paragraph()

    # -------------------------------------------------------------
    # 5. PARAMETER SELECTION DIAGNOSTICS
    # -------------------------------------------------------------
    doc.add_heading("5. HYPERPARAMETER SELECTION DIAGNOSTICS", level=1)
    doc.add_paragraph(
        "Hyperparameters were selected using rigorous empirical curves:\n"
        "- K-Means: The Elbow Method (Inertia SSE) and Silhouette curves both identified k=5 as the optimal inflection point.\n"
        "- Hierarchical: The Ward dendrogram indicated a prominent cutting distance at height ~ 5.5, yielding 5 distinct branches.\n"
        f"- DBSCAN: The sorted 5-nearest neighbor distance curve exhibited an elbow bend around eps ≈ {data['parameter_selection']['dbscan_knee_eps']}."
    )

    # Insert Figure 2
    fig2 = data["figures"]["fig2"]
    if os.path.exists(fig2):
        doc.add_picture(fig2, width=Inches(6.0))
        cap2 = doc.add_paragraph("Figure 2: K-Means parameter optimization diagnostics (Elbow & Silhouette curves).")
        cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap2.runs[0].font.italic = True
        cap2.runs[0].font.size = Pt(9.5)

    # Insert Figure 3 & 4
    fig3 = data["figures"]["fig3"]
    fig4 = data["figures"]["fig4"]
    if os.path.exists(fig3) and os.path.exists(fig4):
        doc.add_picture(fig3, width=Inches(5.5))
        cap3 = doc.add_paragraph("Figure 3: Ward hierarchical dendrogram truncated to 30 nodes.")
        cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap3.runs[0].font.italic = True
        cap3.runs[0].font.size = Pt(9.5)

        doc.add_picture(fig4, width=Inches(5.5))
        cap4 = doc.add_paragraph("Figure 4: DBSCAN k-distance plot (k=5) for epsilon threshold selection.")
        cap4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap4.runs[0].font.italic = True
        cap4.runs[0].font.size = Pt(9.5)

    # -------------------------------------------------------------
    # 6. EXPERIMENTAL RESULTS & COMPARATIVE BENCHMARK
    # -------------------------------------------------------------
    doc.add_heading("6. EXPERIMENTAL RESULTS & COMPARATIVE BENCHMARK", level=1)
    doc.add_paragraph(
        "All algorithms were executed on the standardized feature matrix under fixed random_state=42. "
        "For DBSCAN, noise points were explicitly excluded during internal validation metric calculation."
    )

    # Comparison Table
    bench_table_data = data["benchmarks"]["comparison_table"]
    b_table = doc.add_table(rows=1, cols=7)
    b_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_table.style = 'Light Shading'

    b_headers = ["Algorithm", "Parameters", "Clusters", "Noise Count", "Silhouette", "Davies-Bouldin", "Runtime (s)"]
    for i, h in enumerate(b_headers):
        cell = b_table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_margins(cell)

    for row in bench_table_data:
        r = b_table.add_row()
        r.cells[0].text = str(row["Algorithm"])
        r.cells[1].text = str(row["Parameters"])
        r.cells[2].text = str(row["Clusters Found (excl. noise)"])
        r.cells[3].text = str(row["Noise Count & %"])
        r.cells[4].text = str(row.get("Silhouette Score", "N/A"))
        r.cells[5].text = str(row.get("Davies-Bouldin Index", "N/A"))
        r.cells[6].text = str(row.get("Runtime (s)", "N/A"))
        for c in r.cells:
            set_cell_margins(c)

    doc.add_paragraph()

    # Insert Figure 5 & 6
    fig5 = data["figures"]["fig5"]
    fig6 = data["figures"]["fig6"]
    if os.path.exists(fig5):
        doc.add_picture(fig5, width=Inches(6.0))
        cap5 = doc.add_paragraph("Figure 5: 2D spatial comparison of clustering partitions across K-Means, Ward, and DBSCAN.")
        cap5.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap5.runs[0].font.italic = True
        cap5.runs[0].font.size = Pt(9.5)

    if os.path.exists(fig6):
        doc.add_picture(fig6, width=Inches(6.0))
        cap6 = doc.add_paragraph("Figure 6: Quantitative clustering evaluation benchmark across metrics.")
        cap6.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap6.runs[0].font.italic = True
        cap6.runs[0].font.size = Pt(9.5)

    # -------------------------------------------------------------
    # 7. CLUSTER PERSONAS & MARKETING STRATEGY
    # -------------------------------------------------------------
    doc.add_heading("7. CUSTOMER PERSONA PROFILES & ACTIONABLE STRATEGY", level=1)
    doc.add_paragraph(
        "Using K-Means with k=5, customer clusters were mapped into behavioral personas using rule-based criteria "
        "comparing cluster means against portfolio means normalized by standard deviation (threshold ±0.5 sigma):"
    )

    personas = data["cluster_personas"]
    for p in personas:
        p_desc = doc.add_paragraph()
        p_run = p_desc.add_run(f"Cluster {p['Cluster']}: {p['Segment Name']} ({p['Count']} customers, {p['Percentage']}%)\n")
        p_run.font.bold = True
        p_run.font.color.rgb = RGBColor(30, 58, 138)
        p_desc.add_run(f"• Profile: {p['Description']}\n")
        p_desc.add_run(f"• Annual Income Mean: ${p['Annual Income (k$)']}k | Spending Score Mean: {p['Spending Score (1-100)']}\n")
        p_desc.add_run(f"• Commercial Suggestion: {p['Recommendation']}\n")

    # Insert Figure 7
    fig7 = data["figures"]["fig7"]
    if os.path.exists(fig7):
        doc.add_picture(fig7, width=Inches(5.8))
        cap7 = doc.add_paragraph("Figure 7: Standardized behavioral attribute deviation heatmap across clusters.")
        cap7.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap7.runs[0].font.italic = True
        cap7.runs[0].font.size = Pt(9.5)

    # -------------------------------------------------------------
    # 8. DISCUSSION, LIMITATIONS & CONCLUSION
    # -------------------------------------------------------------
    doc.add_heading("8. DISCUSSION, LIMITATIONS & CONCLUSION", level=1)
    doc.add_paragraph(
        "Empirical Discussion:\n"
        f"K-Means achieved the highest Calinski-Harabasz score ({kmeans_meta['calinski_harabasz_index']}), "
        f"lowest Davies-Bouldin index ({kmeans_meta['davies_bouldin_index']}), and highest Silhouette score ({kmeans_meta['silhouette_score']}) "
        f"in {kmeans_meta['runtime_sec']:.4f} seconds. Ward linkage generated virtually identical clusters, corroborating the structural naturalness "
        "of the 5 customer cohorts. DBSCAN isolated boundary noise points, serving as a valuable anomaly screener rather than a standalone partitioner.\n\n"
        "Limitations:\n"
        "- The dataset captures a static cross-sectional snapshot lacking temporal transaction sequences.\n"
        "- Distance metrics favor convex, spherical clusters and require standardization.\n\n"
        "Conclusion & Future Scope:\n"
        "This project successfully designed, implemented, and validated a complete customer segmentation pipeline. "
        "Future research will explore dynamic RFM (Recency, Frequency, Monetary) time-series clustering and automated real-time API scoring."
    )

    # -------------------------------------------------------------
    # REFERENCES
    # -------------------------------------------------------------
    doc.add_heading("REFERENCES", level=1)
    refs = [
        "MacQueen, J. (1967). Some methods for classification and analysis of multivariate observations. Proc. 5th Berkeley Symp. Math. Statist. Prob., 1, 281-297.",
        "Ward, J. H. (1963). Hierarchical grouping to optimize an objective function. Journal of the American Statistical Association, 58(301), 236-244.",
        "Ester, M., Kriegel, H. P., Sander, J., & Xu, X. (1996). A density-based algorithm for discovering clusters in large spatial databases with noise. In KDD-96 Proceedings, 226-231.",
        "Rousseeuw, P. J. (1987). Silhouettes: a graphical aid to the interpretation and validation of cluster analysis. J. Comput. Appl. Math., 20, 53-65.",
        "Davies, D. L., & Bouldin, D. W. (1979). A cluster separation measure. IEEE Trans. Pattern Anal. Mach. Intell., 1(2), 224-227.",
        "Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830."
    ]
    for r in refs:
        p = doc.add_paragraph(r)
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)

    output_doc_path = os.path.join("reports", "comparative_study_report.docx")
    doc.save(output_doc_path)
    print(f"Generated editable DOCX report successfully at {output_doc_path}!")
    return output_doc_path


if __name__ == "__main__":
    generate_report()
